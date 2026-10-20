"""Opening a workbook for reading, without one wasted pass over a sheet (P08f).

openpyxl opens every worksheet of a read-only workbook by sizing it: it parses the sheet's XML
until the <dimension> element, which Excel writes before <sheetData>. It looks for the end event
of <sheetData> to give up, though, so a sheet that declares no dimension (openpyxl's own
write-only output, and some instrument software) is parsed from end to end at load time, before
the first row is read: for a sheet of a million rows that is a minute spent finding nothing.

open_workbook() loads the workbook exactly as load_workbook(read_only=True, data_only=True) does,
with one difference while it loads: each sheet is sized from start events, which reach
<dimension> just as early and stop at the start of <sheetData>. The dimension found is the same
element, read by the same parser and the same SheetDimension class, so every sheet has the same
size and every row the same width (tests/test_readers_reference.py, and
test_readers_large.py::test_sizing_a_sheet_finds_the_same_dimension). If openpyxl's internals are
not the ones this was written against (3.1), the workbook is loaded unchanged.
"""
from __future__ import annotations

import threading
from contextlib import contextmanager
from pathlib import Path

_lock = threading.Lock()
_depth = 0
_original = None


def _size_from_start_events(ws) -> None:
    """ReadOnlyWorksheet._get_size, stopping at the start of <sheetData>."""
    from openpyxl.worksheet._reader import DATA_TAG, DIMENSION_TAG, iterparse
    from openpyxl.worksheet.dimensions import SheetDimension

    src = ws._get_source()
    try:
        for _event, element in iterparse(src, events=("start",)):
            if element.tag == DIMENSION_TAG:
                ws._min_column, ws._min_row, ws._max_column, ws._max_row = \
                    SheetDimension.from_tree(element).boundaries
                return
            if element.tag == DATA_TAG:
                return                                      # no dimension declared before the data
    finally:
        src.close()


@contextmanager
def _sized_from_start_events():
    global _depth, _original
    try:
        from openpyxl.worksheet import _read_only
        from openpyxl.worksheet._reader import DATA_TAG, DIMENSION_TAG, iterparse  # noqa: F401
        cls = _read_only.ReadOnlyWorksheet
        cls._get_size                                       # noqa: B018 - the hook must exist
    except (ImportError, AttributeError):
        yield
        return
    with _lock:
        if _depth == 0:
            _original = cls._get_size
            cls._get_size = _size_from_start_events
        _depth += 1
    try:
        yield
    finally:
        with _lock:
            _depth -= 1
            if _depth == 0:
                cls._get_size = _original


def open_workbook(path: str | Path):
    """load_workbook(path, read_only=True, data_only=True), sizing each sheet without a full pass."""
    from openpyxl import load_workbook

    with _sized_from_start_events():
        return load_workbook(path, read_only=True, data_only=True)
