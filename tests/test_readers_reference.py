"""The reader-equivalence gate (P08f): the streaming reader returns what the whole-grid reader did.

tests/fixtures/reader_reference/ was captured by scripts/capture_reader_reference.py from the
reader as it was before the streaming rewrite. For every shipped public instrument file, every
BioLogic text export read through the generic engine and every case of the synthetic corpus
(tests/reader_corpus.py), the reader is run again here and compared field by field:

- the outcome (a table, or an exception of the same type with the same code and message);
- every array (t, i, e, step_id, cycle, q_charge, q_discharge, dq, temperature): same dtype,
  same shape, identical values with NaN equal to NaN (np.array_equal(equal_nan=True) against
  the stored arrays of the real files, and a NaN-canonical SHA-256 for every array);
- step_kind, the column names, the import report and the complete findings (code, severity,
  message and detail, in order);
- for the real files, the profile the registry identifies, and the proposer's output.

The corpus is also read with the streaming block and chunk sizes set to a few bytes and a few
rows, which puts a block or chunk boundary inside every line, every multibyte character and every
CR LF pair; the answer must not change.
"""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import numpy as np
import pytest

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "reader_reference"


def _capture_module():
    spec = importlib.util.spec_from_file_location("capture_reader_reference",
                                                  ROOT / "scripts" / "capture_reader_reference.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


cap = _capture_module()
REF = json.loads((FIXTURE / "reference.json").read_text(encoding="utf-8"))
ARRAYS = np.load(FIXTURE / "arrays.npz")


@pytest.fixture(autouse=True)
def _no_saved_profiles(tmp_path, monkeypatch):
    monkeypatch.setenv("ELECTROCHEM_PROFILES", str(tmp_path / "profiles"))


def same(new: dict, old: dict, key: str, arrays: dict, what: str) -> None:
    """Assert that one captured outcome equals its reference."""
    assert new["kind"] == old["kind"], f"{what}: {new.get('message', '')[:300]} vs {old.get('message', '')[:300]}"
    if new["kind"] == "error":
        assert (new["type"], new["code"]) == (old["type"], old["code"]), what
        assert new["message"] == old["message"], what
        return
    if new["kind"] == "proposal":
        for field in old:
            if field in ("layout_rows", "layout_rows_sha256"):
                continue                       # compared by the caller where the window covers the file
            assert new[field] == old[field], f"{what}: proposal field {field}"
        return
    for field in ("source_name", "source_sha256", "profile", "n", "columns", "step_kind", "report"):
        assert new[field] == old[field], f"{what}: {field}"
    assert new["findings"] == old["findings"], f"{what}: findings (code, severity, message, detail)"
    assert new["arrays"].keys() == old["arrays"].keys()
    for name, meta in old["arrays"].items():
        assert (new["arrays"][name] is None) == (meta is None), f"{what}: {name} present"
        if meta is None:
            continue
        assert new["arrays"][name] == meta, f"{what}: {name} differs (dtype, shape or values)"
        stored = f"{key}::{name}"
        if stored in ARRAYS.files:
            a, b = arrays[stored], ARRAYS[stored]
            assert a.dtype == b.dtype and a.shape == b.shape, f"{what}: {name}"
            assert np.array_equal(a, b, equal_nan=True), f"{what}: {name}"


# ----------------------------------------------------------------- real files
PUBLIC = sorted(REF["public"])


@pytest.fixture(scope="module")
def registry(tmp_path_factory):
    from electrochem.readers import Registry
    return Registry(tmp_path_factory.mktemp("noprofiles") / "profiles")


@pytest.mark.parametrize("rel", PUBLIC)
def test_a_shipped_file_reads_exactly_as_before(rel, registry):
    from electrochem.readers import BUILTIN
    old = REF["public"][rel]
    path = ROOT / rel
    assert cap.sha256_file(path) == old["sha256"], "the shipped file changed since the reference was taken"
    arrays: dict = {}
    if "identified" not in old:                                   # a BioLogic export, generic engine
        from tests.test_readers_differential import BIOLOGIC_MPT
        same(cap.capture_read(path, BIOLOGIC_MPT, arrays, f"{rel}::read"), old["read"], f"{rel}::read", arrays, rel)
        return
    try:
        prof = registry.identify(path, list(BUILTIN.values()))
        assert old["identified"] == prof.name, rel
    except Exception as exc:
        prof = None
        assert old["identified"] == cap.dump_error(exc), rel
    assert ("read" in old) == (prof is not None)
    if prof is not None:
        same(cap.capture_read(path, prof, arrays, f"{rel}::read"), old["read"], f"{rel}::read", arrays, rel)
        if "read_unsigned" in old:
            answered = prof.model_copy(update={"conventions": prof.conventions.model_copy(
                update={"current_sign": "unsigned"})})
            same(cap.capture_read(path, answered, arrays, f"{rel}::read_unsigned"), old["read_unsigned"],
                 f"{rel}::read_unsigned", arrays, rel + " (unsigned)")
    new = cap.capture_propose(path)
    same(new, old["propose"], "", {}, rel + " (proposal)")
    if new["kind"] == "proposal" and path.suffix.lower() not in (".xlsx", ".xlsm", ".xls"):
        # every shipped text file lies inside the proposer's detection window, so the rows
        # detect_layout returns are the whole table, as before
        assert (new["layout_rows"], new["layout_rows_sha256"]) == \
            (old["propose"]["layout_rows"], old["propose"]["layout_rows_sha256"]), rel


# ------------------------------------------------------------------ the corpus
@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    from tests.reader_corpus import write_all
    return write_all(tmp_path_factory.mktemp("corpus"))


CORPUS = sorted(REF["corpus"])


def _corpus_case(corpus, name):
    """The generated input, its profile and its reference. A generated workbook may differ in its bytes
    from the one captured while holding the same cells: openpyxl's XML writer follows ElementTree, which
    writes `<x />` under Python 3.12 (the installers) and `<x/>` under 3.13, and orders namespace
    declarations differently. Then the cells must match the captured content digest, and the reader's
    own source hash is compared with the bytes it actually read instead of the captured bytes; every
    other part of the outcome (arrays, findings, report, step kinds) is compared as captured."""
    import copy
    import hashlib
    hit = [(p, prof) for n, p, prof in corpus if n == name]
    assert hit, f"{name} is in the reference but no longer in tests/reader_corpus.py"
    path, prof = hit[0]
    old = REF["corpus"][name]
    digest = cap.sha256_file(path)
    if digest != old["sha256"]:
        assert "content_sha256" in old and cap.workbook_digest(path) == old["content_sha256"], \
            f"{name}: the generator writes other content than at capture"
        old = copy.deepcopy(old)
        if old["read"]["kind"] == "table":
            old["read"]["source_sha256"] = digest
    assert hashlib.sha256(cap.canonical(prof.model_dump(mode="json")).encode()).hexdigest() == old["profile_sha256"]
    return path, prof, old


def test_a_generated_workbook_is_recognised_by_its_cells_when_its_bytes_differ(corpus, tmp_path):
    """The relaxation above admits other bytes only with the same cells: a workbook with one cell changed
    fails the content check, and the same workbook re-saved with other XML formatting passes it."""
    import shutil
    import zipfile
    name = next(n for n in CORPUS if n.endswith(".xlsx") and "content_sha256" in REF["corpus"][n])
    src = next(p for n, p, _ in corpus if n == name)
    old = REF["corpus"][name]
    assert cap.workbook_digest(src) == old["content_sha256"]
    respaced = tmp_path / "respaced.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(respaced, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename.endswith(".xml"):
                data = data.replace(b'"/>', b'" />')                 # what Python 3.12's writer produces
            zout.writestr(info.filename, data)
    assert cap.sha256_file(respaced) != old["sha256"] and cap.workbook_digest(respaced) == old["content_sha256"]
    from openpyxl import load_workbook
    changed = tmp_path / "changed.xlsx"
    shutil.copy(src, changed)
    wb = load_workbook(changed)
    ws = wb.worksheets[-1]
    cell = next(c for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float)))
    cell.value = cell.value + 1
    wb.save(changed)
    assert cap.workbook_digest(changed) != old["content_sha256"]


def test_the_corpus_is_the_one_captured(corpus):
    assert sorted(n for n, _, _ in corpus) == CORPUS


@pytest.mark.parametrize("name", CORPUS)
def test_a_corpus_case_reads_exactly_as_before(corpus, name):
    path, prof, old = _corpus_case(corpus, name)
    same(cap.capture_read(path, prof, None, ""), old["read"], "", {}, name)
    if "propose" in old:
        new = cap.capture_propose(path)
        same(new, old["propose"], "", {}, name + " (proposal)")
        if new["kind"] == "proposal" and path.suffix.lower() != ".xlsx":
            assert (new["layout_rows"], new["layout_rows_sha256"]) == \
                (old["propose"]["layout_rows"], old["propose"]["layout_rows_sha256"]), name


@pytest.mark.parametrize("block,chunk", [(5, 3), (4099, 1009)])
def test_block_and_chunk_boundaries_change_nothing(corpus, monkeypatch, block, chunk):
    """Every text case with blocks of a few bytes and chunks of a few rows (the long cases with
    blocks just over 4 KiB), every workbook case with chunks of a few rows."""
    from electrochem.readers import reader
    monkeypatch.setattr(reader, "CHUNK_ROWS", chunk)
    monkeypatch.setattr(reader, "_SPLIT_BELOW", 2)
    for name in CORPUS:
        path, prof, old = _corpus_case(corpus, name)
        big = name.startswith("big_")
        if (block < 1000) == big:
            continue
        monkeypatch.setattr(reader, "BLOCK_BYTES", block)
        same(cap.capture_read(path, prof, None, ""), old["read"], "", {}, f"{name} (block {block}, chunk {chunk})")


def test_the_progress_callback_sees_every_row_once(corpus):
    from electrochem.readers.reader import read
    path, prof, old = _corpus_case(corpus, "big_mixed_crlf.csv")
    calls = []
    tab = read(path, prof, progress=lambda rows, done, total: calls.append((rows, done, total)))
    assert len(calls) >= 2
    rows, done, total = zip(*calls)
    assert list(rows) == sorted(rows) and list(done) == sorted(done)
    assert rows[-1] == tab.n and done[-1] == total[-1] == path.stat().st_size
    assert set(total) == {path.stat().st_size}


@pytest.mark.parametrize("over", [{"header_row": 0}, {"header_row": -2}, {"unit_row": -1},
                                  {"skip_rows_after_header": -1}])
def test_a_layout_that_counts_rows_from_the_end_is_refused(tmp_path, over):
    """The whole-grid reader indexed its list of rows with these, so a header row of 0 took the last
    line of the file as the header. No layout means that; the streaming reader refuses it by name."""
    from electrochem.readers.reader import ImportFailure, read
    from tests.reader_corpus import SIMPLE, _simple_text, _with_layout
    p = tmp_path / "x.csv"
    p.write_text(_simple_text())
    with pytest.raises(ImportFailure, match="invalid_layout"):
        read(p, _with_layout(SIMPLE, **over))
    assert read(p, _with_layout(SIMPLE, header_row=0, header_contains=("Test_Time",))).n == 40


@pytest.mark.parametrize("decimal", [".", ","])
def test_the_chunk_conversion_equals_the_token_by_token_conversion(decimal):
    """_numbers (a whole chunk through float(), with the fallbacks) against _to_float (the
    original, one token at a time) on random tokens built from everything the fast path has to
    rule out: whitespace inside and around, U+001C to U+001F, Unicode minus, digits and spaces,
    'nan', 'inf', underscores, d/D/e/E exponents, both decimal marks, grouping, NUL, empty."""
    import random
    from electrochem.readers.reader import _numbers, _to_float
    rng = random.Random(8101)
    pieces = ["0", "1", "7", "42", ".", ",", "-", "+", "e", "E", "d", "D", "5", "e-3", "D+02", " ", "\t",
              "\x1f", "\x1c", "−", " ", "５", "٣", "nan", "NaN", "inf", "Infinity", "_",
              "\x00", "x", "1,234", "1.234", "12.5", "0.001"]
    common = ["0.5", "-1.25", "3", "1e-3", "2.5E+01", " 4.5", "6 ", "7,5", "1.234,5", "1,234.5", "2D-03"]
    for trial in range(300):
        n = rng.choice([1, 2, 3, 17, 64, 1500, 3000])
        toks = []
        for _ in range(n):
            if rng.random() < 0.85:
                toks.append(rng.choice(common))
            else:
                toks.append("".join(rng.choice(pieces) for _ in range(rng.randint(0, 4))))
        a, bad_a = _numbers(tuple(toks), decimal)
        b, bad_b = _to_float(toks, decimal)
        assert bad_a == bad_b, (trial, toks[:20])
        assert a.dtype == b.dtype and np.array_equal(a, b, equal_nan=True), (trial, toks[:20])
        assert np.array_equal(np.signbit(a), np.signbit(b))
    clean = tuple(rng.choice(common[:5]) for _ in range(5000))
    from electrochem.readers.reader import _numbers_fast
    assert _numbers_fast(clean, decimal) is not None                     # a clean chunk takes the fast path
