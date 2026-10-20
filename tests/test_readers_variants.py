"""Export variants (P08e). Arbin, Landt and the rest let the user choose, at export time, which
columns to write, in what order, in which units, and under which sheet name; two exports of the
same test from the same software rarely share a header. The reader must treat all of them as the
same instrument and give the same numbers: a metamorphic property (the transform's effect on the
output is known exactly) that needs no reference value.

Three layers are exercised: a built-in profile (Arbin, Landt) on re-exports of a shipped real
file; a profile learned once from an unknown instrument, on later exports with other columns
and units; and the review page, which starts from the saved profile when a required column was
renamed, so the person confirms instead of answering again."""
import csv
import json
from pathlib import Path

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook

from electrochem import workflow
from electrochem.cycling import Settings
from electrochem.readers import Conventions, Registry, confirm

ROOT = Path(__file__).resolve().parents[1]
ARBIN = ROOT / "data" / "public" / "arbin" / "bs542_004_gr_li_50ua_50mv_1v_191020_Channel_11.xlsx"
LANDT = ROOT / "data" / "public" / "landt" / "NJK_CC_156_C30_C_x30.xlsx"


@pytest.fixture
def profiles(tmp_path, monkeypatch):
    monkeypatch.setenv("ELECTROCHEM_PROFILES", str(tmp_path / "profiles"))
    return tmp_path / "profiles"


def _sheet(path, name, limit=None):
    wb = load_workbook(path, read_only=True)
    rows = []
    for k, r in enumerate(wb[name].iter_rows(values_only=True)):
        rows.append(list(r))
        if limit and k >= limit:
            break
    wb.close()
    return rows[0], rows[1:]


def _write_xlsx(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def _write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    return path


def _variant(header, rows, keep=None, rename=None, scale=None, extra=None, order=None):
    """Re-export: keep a subset, rename (unit changes), scale the values to the new unit, add
    columns, reorder."""
    rename, scale, extra = rename or {}, scale or {}, extra or []
    cols = [h for h in header if keep is None or h in keep]
    idx = {h: header.index(h) for h in cols}
    out_header = [rename.get(h, h) for h in cols] + [name for name, _ in extra]
    out_rows = []
    for k, r in enumerate(rows):
        row = []
        for h in cols:
            v = r[idx[h]]
            if h in scale and isinstance(v, (int, float)):
                v = v * scale[h]
            row.append(v)
        row += [fn(k) for _, fn in extra]
        out_rows.append(row)
    if order:
        perm = [out_header.index(h) for h in order] + [k for k, h in enumerate(out_header) if h not in order]
        out_header = [out_header[k] for k in perm]
        out_rows = [[r[k] for k in perm] for r in out_rows]
    return out_header, out_rows


def _numbers(L):
    """Branch capacities on the instrument-counter basis where the file carries a counter, and the
    trapezoid of the sampled current always (M01: the counter is preferred when present, so a
    re-export without the counter columns legitimately changes the primary basis, not the data)."""
    run = workflow.analyse_file(L, Settings(charge="ox", terminal_voltage=True))
    q = {b.direction: (b.charge_mAh, b.charge_trapezoid_mAh, b.charge_method) for b in run.an.branches}
    return q, L.tab


def _same(a, b, atol=1e-9):
    assert np.allclose(a.t, b.t, rtol=0, atol=1e-6) and np.allclose(a.i, b.i, rtol=1e-12, atol=atol) \
        and np.allclose(a.e, b.e, rtol=1e-12, atol=atol)


# ---------------------------------------------------------------- Arbin
@pytest.fixture(scope="module")
def arbin():
    header, rows = _sheet(ARBIN, "Channel_11_1")
    rows = [r for r in rows if r[1] is not None]
    for r in rows:
        r[0] = str(r[0])                                  # the wall clock as text, as the CSV export writes it
    return header, rows


@pytest.fixture(scope="module")
def arbin_reference(arbin, tmp_path_factory):
    header, rows = arbin
    p = _write_xlsx(tmp_path_factory.mktemp("ref") / "ref.xlsx", [("Channel_11_1", header, rows)])
    L = workflow.open_file(p)
    assert L.proposal is None and L.meta["profile"] == "arbin-xlsx"
    q, tab = _numbers(L)
    return q, tab


@pytest.mark.parametrize("case", ["subset", "reorder_extra", "units", "sheet_names", "csv_subset", "csv_units"])
def test_an_arbin_export_with_other_columns_order_units_or_sheets_is_the_same_test(arbin, arbin_reference, tmp_path, profiles, case):
    header, rows = arbin
    q0, tab0 = arbin_reference
    extra = [("Aux_Temperature_1(C)", lambda k: 25.0 + 0.01 * (k % 7)), ("ACR(Ohm)", lambda k: 0.5),
             ("Is_FC_Data", lambda k: 0), ("Aux_Voltage_1(V)", lambda k: 0.001 * k)]
    if case == "subset":                                   # only what Arbin's dialog ticks by default
        h, r = _variant(header, rows, keep=["Test_Time(s)", "Step_Index", "Cycle_Index", "Voltage(V)", "Current(A)"])
        p = _write_xlsx(tmp_path / "subset.xlsx", [("Channel_11_1", h, r)])
    elif case == "reorder_extra":
        h, r = _variant(header, rows, extra=extra,
                        order=["Voltage(V)", "Current(A)", "Aux_Voltage_1(V)", "Cycle_Index", "Test_Time(s)"])
        p = _write_xlsx(tmp_path / "reorder.xlsx", [("Channel_11_1", h, r)])
    elif case == "units":                                  # mA, mAh and minutes chosen at export time
        h, r = _variant(header, rows, rename={"Current(A)": "Current(mA)", "Charge_Capacity(Ah)": "Charge_Capacity(mAh)",
                                              "Discharge_Capacity(Ah)": "Discharge_Capacity(mAh)",
                                              "Test_Time(s)": "Test_Time(min)"},
                        scale={"Current(A)": 1e3, "Charge_Capacity(Ah)": 1e3, "Discharge_Capacity(Ah)": 1e3,
                               "Test_Time(s)": 1 / 60})
        p = _write_xlsx(tmp_path / "units.xlsx", [("Channel_11_1", h, r)])
    elif case == "sheet_names":                            # another channel naming, a statistics sheet, the info sheet
        h, r = _variant(header, rows, extra=extra[:1])
        stats_h = ["Cycle_Index", "Charge_Capacity(Ah)", "Discharge_Capacity(Ah)", "Coulombic_Efficiency(%)"]
        p = _write_xlsx(tmp_path / "sheets.xlsx", [("Global_Info", ["Channel", "Start DateTime"], [[11, "2020-10-19"]]),
                                                   ("Channel_1-1", h, r),
                                                   ("Statistics_1-1", stats_h, [[1, 0.001, 0.0009, 90.0]])])
    elif case == "csv_subset":
        h, r = _variant(header, rows, keep=["Data_Point", "Test_Time(s)", "Current(A)", "Voltage(V)", "Cycle_Index",
                                            "Step_Index", "Charge_Capacity(Ah)", "Discharge_Capacity(Ah)"])
        p = _write_csv(tmp_path / "subset.csv", h, r)
    else:
        h, r = _variant(header, rows, rename={"Current(A)": "Current(mA)", "Voltage(V)": "Voltage(mV)"},
                        scale={"Current(A)": 1e3, "Voltage(V)": 1e3}, extra=extra)
        p = _write_csv(tmp_path / "units.csv", h, r)
    L = workflow.open_file(p)
    assert L.proposal is None, "the same instrument must not be asked about again"
    assert L.meta["profile"] in ("arbin-xlsx", "arbin-csv")
    q, tab = _numbers(L)
    _same(tab, tab0, atol=1e-9 if case != "csv_units" else 1e-6)
    counter = "Charge_Capacity(mAh)" in h or "Charge_Capacity(Ah)" in h
    for d in ("charge", "discharge"):
        if d in q0:
            assert q[d][1] == pytest.approx(q0[d][1], rel=1e-6), (case, d, "trapezoid")
            if counter:
                assert q[d][0] == pytest.approx(q0[d][0], rel=1e-6) and q[d][2] == q0[d][2], (case, d, "counter")
            else:
                assert q[d][0] == q[d][1], "without the counter columns the trapezoid is the capacity"
    if counter:
        assert L.meta["import_report"].get("counter_vs_integral", {}).get("compared", True) is not False


# ---------------------------------------------------------------- Landt
@pytest.fixture(scope="module")
def landt():
    header, rows = _sheet(LANDT, "Record", limit=4000)
    return header, [r for r in rows if r[1] is not None]


@pytest.fixture(scope="module")
def landt_reference(landt, tmp_path_factory):
    header, rows = landt
    p = _write_xlsx(tmp_path_factory.mktemp("ref") / "ref.xlsx", [("Record", header, rows)])
    L = workflow.open_file(p)
    assert L.proposal is None and L.meta["profile"] == "landt-xlsx"
    return _numbers(L)


@pytest.mark.parametrize("case", ["units", "subset_extra", "sheet_name_and_order"])
def test_a_landt_export_in_other_units_columns_or_sheet_is_the_same_test(landt, landt_reference, tmp_path, profiles, case):
    header, rows = landt
    q0, tab0 = landt_reference
    if case == "units":                                    # A, mV, Ah and seconds chosen in LANDdt's export dialog
        h, r = _variant(header, rows, rename={"Current/mA": "Current/A", "Voltage/V": "Voltage/mV",
                                              "Capacity/mAh": "Capacity/Ah", "TestTime/h": "TestTime/s",
                                              "StepTime/h": "StepTime/s"},
                        scale={"Current/mA": 1e-3, "Voltage/V": 1e3, "Capacity/mAh": 1e-3, "TestTime/h": 3600,
                               "StepTime/h": 3600})
        p = _write_xlsx(tmp_path / "units.xlsx", [("Record", h, r)])
    elif case == "subset_extra":
        h, r = _variant(header, rows, keep=["TestTime/h", "Voltage/V", "Current/mA", "State"],
                        extra=[("Energy/mWh", lambda k: 0.0), ("SOC/DOD/%", lambda k: 50.0), ("Temperature/C", lambda k: 25.0)])
        p = _write_xlsx(tmp_path / "subset.xlsx", [("Record", h, r)])
    else:
        h, r = _variant(header, rows, order=["State", "Current/mA", "Voltage/V", "TestTime/h"])
        p = _write_xlsx(tmp_path / "sheet.xlsx", [("Sheet1", h, r),
                                                  ("Cycle", ["Index", "Charge-Cap/mAh", "Discharge-Cap/mAh"], [[1, 0.0176, 0.0256]])])
    L = workflow.open_file(p)
    assert L.proposal is None and L.meta["profile"] == "landt-xlsx"
    q, tab = _numbers(L)
    _same(tab, tab0, atol=1e-9)
    for d in q0:
        assert q[d] == pytest.approx(q0[d], rel=1e-6), (case, d)        # Landt: no counter bound, so one basis


# ---------------------------------------------------------------- a learned instrument
DEFAULT_STRANGE = ("Elapsed Time(s)", "Amps(A)", "Volts(V)", "Cycle Index", "Step Index")


def _strange(path, header=DEFAULT_STRANGE, i_scale=1.0, extra=(), order=None):
    """An export from an invented instrument: the five default columns, optionally renamed (header,
    positionally), with the current scaled to a new unit, extra columns, and a new column order."""
    rows = []
    for k in range(120):
        i = (0.001 if k < 60 else -0.001) * i_scale
        e = 3.5 + 0.001 * k
        values = [k * 60, i, e, 1, 1 if k < 60 else 2]
        row = dict(zip(header, values))
        for n, fn in extra:
            row[n] = fn(k)
        rows.append(row)
    names = list(header) + [n for n, _ in extra]
    if order:
        names = list(order) + [n for n in names if n not in order]
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(names)
        for row in rows:
            w.writerow([row[n] for n in names])
    return path


def _learn(tmp_path, profiles):
    p = _strange(tmp_path / "first.csv")
    L = workflow.open_file(p)
    assert L.proposal is not None and L.proposal.prior is None
    conv = Conventions(current_sign="charge_positive", direction_source="current_sign", capacity_semantics="none",
                       cycle_base=1, evidence="confirmed in the test against the writer's own numbers")
    confirm(L.proposal, name="strange-csv", instrument="Strange cycler 3000", conventions=conv, registry=Registry(profiles))
    return workflow.analyse_file(workflow.open_file(p), Settings(charge="ox", terminal_voltage=True))


def test_a_learned_instrument_is_recognised_with_other_columns_order_and_units(tmp_path, profiles):
    """Learned once; later exports with extra columns, another order or another unit open with no questions."""
    run0 = _learn(tmp_path, profiles)
    q0 = {b.direction: b.charge_mAh for b in run0.an.branches}
    later = _strange(tmp_path / "later.csv", extra=(("Temp(C)", lambda k: 25.0), ("dV/dt(V/s)", lambda k: 0.0)),
                     order=("Volts(V)", "Temp(C)", "Amps(A)"))
    L = workflow.open_file(later)
    assert L.proposal is None and L.meta["profile"] == "strange-csv"
    q = {b.direction: b.charge_mAh for b in workflow.analyse_file(L, Settings(charge="ox", terminal_voltage=True)).an.branches}
    assert q == pytest.approx(q0)
    ma = _strange(tmp_path / "milliamps.csv", header=("Elapsed Time(s)", "Amps(mA)", "Volts(V)", "Cycle Index", "Step Index"),
                  i_scale=1e3)
    L = workflow.open_file(ma)
    assert L.proposal is None and L.meta["profile"] == "strange-csv"
    q = {b.direction: b.charge_mAh for b in workflow.analyse_file(L, Settings(charge="ox", terminal_voltage=True)).an.branches}
    assert q == pytest.approx(q0), "the unit in the header, not the unit of the first file, decides"
    saved = json.loads((profiles / "strange-csv.json").read_text())
    assert saved["profile"]["bindings"]["current_ampere"]["source_unit"] is None   # a stated unit is not frozen


def test_a_renamed_column_asks_again_but_starts_from_the_saved_profile(tmp_path, profiles):
    """A required column renamed between exports cannot be matched silently; the proposal then carries
    the saved profile of the instrument whose other columns match, so the person confirms rather
    than answers from nothing."""
    _learn(tmp_path, profiles)
    renamed = _strange(tmp_path / "renamed.csv", header=("Total Elapsed(s)", "Amps(A)", "Volts(V)", "Cycle Index", "Step Index"))
    L = workflow.open_file(renamed)
    assert L.proposal is not None
    assert L.proposal.prior is not None and L.proposal.prior.name == "strange-csv"
    assert {"Amps(A)", "Volts(V)", "Cycle Index", "Step Index"} <= set(L.proposal.prior_shared)
    assert any("strange-csv" in q for q in L.proposal.questions)
    other = tmp_path / "other.csv"
    other.write_text("Zeit(s),Strom(A),Spannung(V)\n" + "\n".join(f"{k},0.001,3.5" for k in range(10)) + "\n")
    assert workflow.open_file(other).proposal.prior is None             # nothing in common: no prior offered


def test_the_registry_never_lets_two_saved_profiles_claim_silently(tmp_path, profiles):
    _learn(tmp_path, profiles)
    L = workflow.open_file(_strange(tmp_path / "again.csv"))
    conv = Conventions(current_sign="discharge_positive", direction_source="current_sign", capacity_semantics="none",
                       cycle_base=1, evidence="a second person decided the opposite sign")
    reg = Registry(profiles)
    prop = workflow.open_file(_strange(tmp_path / "b.csv", extra=(("Extra", lambda k: 1),)))
    assert prop.proposal is None                                          # claimed by strange-csv
    # a second saved profile binding the same columns with another sign: the file is ambiguous and is asked about
    from electrochem.readers import propose_mapping
    p2 = propose_mapping(tmp_path / "b.csv")
    confirm(p2, name="strange-b", instrument="Strange cycler, other sign", conventions=conv, registry=reg)
    L = workflow.open_file(_strange(tmp_path / "c.csv", extra=(("Other", lambda k: 2),)))
    assert L.proposal is not None, "two saved profiles claim the file: the choice is the person's"


def test_two_channels_in_one_workbook_are_refused_not_merged(arbin, tmp_path, profiles):
    """Arbin can write several channels into one workbook, each sheet with the same header. Read as one
    record they would be one cell with twice the capacity; the clock restarting is the giveaway."""
    header, rows = arbin
    half = rows[: len(rows) // 2]
    p = _write_xlsx(tmp_path / "two_channels.xlsx", [("Channel_1_1", header, half), ("Channel_2_1", header, half)])
    from electrochem.readers import ImportFailure
    with pytest.raises(ImportFailure) as exc:
        workflow.open_file(p)
    assert exc.value.code == "sheets_not_continuous" and "Channel_2_1" in str(exc.value)
    # a long record split across sheets, the clock running on, is still read as one
    a, b = rows[: len(rows) // 2], rows[len(rows) // 2:]
    q = _write_xlsx(tmp_path / "continued.xlsx", [("Channel_1_1", header, a), ("Channel_1_1_Wb_1", header, b)])
    L = workflow.open_file(q)
    assert L.proposal is None and L.tab.t.size == len(rows)
    assert any(f["code"] == "sheets_concatenated" for f in L.meta["import_findings"])
