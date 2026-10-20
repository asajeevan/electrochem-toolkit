"""One known experiment, written out in three vendor dialects.

Ground truth, exact by construction:
    rest 300 s, charge +1 mA for 3600 s, rest 300 s, discharge 1 mA for 1800 s
    charge capacity   = 1.0 mAh   = 1e-3 Ah
    discharge capacity = 0.5 mAh  = 5e-4 Ah
    coulombic efficiency = 50 %

The dialects differ in everything that normally goes wrong: units, signed or
unsigned current, where direction lives, capacity semantics, decimal separator,
clock-format or numeric time, file kind and sheet layout. Any reader that gets
the same answer from all three has understood the conventions rather than the
file.

Each writer reproduces the layout of a real export read in P08d (column names,
units, preamble, sheet structure, counter behaviour), so the synthetic known
answer and the real files exercise the same profile: Arbin MITS Pro CSV (the
column names of galv-harvester's arbin.csv and navani's Arbin workbook),
Neware's record sheet (PyProBE's sample workbook, plus the Step Type and
DataPoint columns navani's reader documents), and Maccor's export with the
12-line header block (cellpy's maccor_002.txt and maccor_two.txt).
"""
from __future__ import annotations

from pathlib import Path

SEGMENTS = [          # (kind, current_mA, duration_s, v_start, v_end)
    ("rest", 0.0, 300.0, 3.00, 3.00),
    ("charge", 1.0, 3600.0, 3.00, 4.00),
    ("rest", 0.0, 300.0, 4.00, 3.98),
    ("discharge", -1.0, 1800.0, 3.98, 3.50),
]
DT = 60.0
TRUE_Q_CHARGE_AH = 1.0e-3
TRUE_Q_DISCHARGE_AH = 5.0e-4


def ground_truth() -> list[dict]:
    """Rows with everything any dialect might need, in SI plus convenience units."""
    rows: list[dict] = []
    t = 0.0
    qc = qd = 0.0           # test-cumulative, Ah
    sqc = sqd = 0.0         # step-cumulative, Ah
    for step, (kind, i_ma, dur, v0, v1) in enumerate(SEGMENTS):
        sqc = sqd = 0.0
        n = int(dur / DT) + 1
        for k in range(n):
            frac = k / (n - 1)
            if k > 0:                       # charge accumulated over the last DT
                inc = abs(i_ma) * 1e-3 * DT / 3600.0
                if i_ma > 0:
                    qc += inc; sqc += inc
                elif i_ma < 0:
                    qd += inc; sqd += inc
            rows.append(
                dict(
                    t=t, step=step, cycle=1 if step < 4 else 2, kind=kind,
                    i_ma=i_ma, v=v0 + frac * (v1 - v0),
                    qc_ah=qc, qd_ah=qd, sqc_ah=sqc, sqd_ah=sqd,
                )
            )
            if k < n - 1:
                t += DT
        t += DT
    return rows


def _f(x: float, decimal: str = ".", nd: int = 9) -> str:
    s = f"{x:.{nd}f}"
    return s.replace(".", ",") if decimal == "," else s


def write_arbin_like(path: Path) -> Path:
    """Signed current in A, test-cumulative capacity in Ah, comma-separated."""
    head = ["Data_Point", "Test_Time(s)", "Step_Index", "Cycle_Index",
            "Current(A)", "Voltage(V)", "Charge_Capacity(Ah)", "Discharge_Capacity(Ah)"]
    out = [",".join(head)]
    for k, r in enumerate(ground_truth(), 1):
        out.append(",".join([
            str(k), _f(r["t"], nd=3), str(r["step"] + 1), str(r["cycle"]),
            _f(r["i_ma"] * 1e-3), _f(r["v"], nd=6), _f(r["qc_ah"]), _f(r["qd_ah"]),
        ]))
    path.write_text("\n".join(out) + "\n", encoding="utf-8", newline="\n")
    return path


def write_neware_like(path: Path, sheets: int = 2) -> Path:
    """Neware's Excel record sheet: signed current in mA, step-cumulative mAh counters that restart
    at every step, Total Time as an Excel time cell (read back by openpyxl as a time span), a
    statistics sheet before the records that must be ignored, and the records split across sheets
    that repeat the header (the split names are this fixture's; the real file read had one)."""
    import datetime as dt

    from openpyxl import Workbook

    rows = ground_truth()
    head = ["DataPoint", "Date", "Cycle Index", "Step Index", "Step Type", "Time", "Total Time",
            "Current(mA)", "Voltage(V)", "Capacity(mAh)", "Chg. Cap.(mAh)", "DChg. Cap.(mAh)"]
    status = {"rest": "Rest", "charge": "CC Chg", "discharge": "CC DChg"}
    start = dt.datetime(2026, 2, 1, 9, 0, 0)
    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("cycle")
    summary.append(["Cycle Index", "Chg. Cap.(mAh)", "DChg. Cap.(mAh)"])
    summary.append([1, 1.0, 0.5])
    per = (len(rows) + sheets - 1) // sheets
    step_start = {}
    for r in rows:
        step_start.setdefault(r["step"], r["t"])
    for s in range(sheets):
        ws = wb.create_sheet("record" if s == 0 else f"record_{s + 1}")
        ws.append(head)
        for k, r in enumerate(rows[s * per:(s + 1) * per], 1 + s * per):
            q = (r["sqc_ah"] + r["sqd_ah"]) * 1e3
            ws.append([k, start + dt.timedelta(seconds=r["t"]), r["cycle"], r["step"] + 1, status[r["kind"]],
                       dt.timedelta(seconds=r["t"] - step_start[r["step"]]), dt.timedelta(seconds=r["t"]),
                       r["i_ma"], r["v"], q, r["sqc_ah"] * 1e3, r["sqd_ah"] * 1e3])
    wb.save(path)
    return path


def write_maccor_like(path: Path, decimal: str = ",") -> Path:
    """Maccor's export with a 12-line header block: unsigned Current [A] with the direction in Md,
    one unsigned Cap. [Ah] counter that restarts at every step, TestTime in minutes with no unit in
    its name, a blank line after the header, decimal comma (KIT's maccor_two.txt; maccor_002.txt
    is the same layout with a decimal point)."""
    head = ["Rec", "Cycle P", "Cycle C", "Step", "TestTime", "StepTime", "Cap. [Ah]", "Ener. [Wh]",
            "Current [A]", "Voltage [V]", "Md", "ES", "DPT Time"]
    out = ["Name:\tDEMO-001", "Channel:\t1", "Tester\tMaccor #1", "Procedure\tdemo.000",
           "Description\tsynthetic fixture", "Started\t02/01/2026 09:00:00", f"SCap\t{_f(0.001, decimal, 3)}",
           f"C Rate\t{_f(1.0, decimal, 3)}", f"Ch. max. V\t{_f(4.2, decimal, 2)} V",
           f"DisCh. min. V\t{_f(2.5, decimal, 2)} V", f"Ch. max. I\t{_f(0.01, decimal, 3)} A",
           f"DisCh. max. I\t{_f(0.01, decimal, 3)} A", ""]
    out.append("\t".join(head))
    out.append("")
    state = {"rest": "R", "charge": "C", "discharge": "D"}
    step_start = {}
    for r in ground_truth():
        step_start.setdefault(r["step"], r["t"])
    for k, r in enumerate(ground_truth(), 1):
        q = r["sqc_ah"] + r["sqd_ah"]
        out.append("\t".join([
            str(k), "0", str(r["cycle"]), str(r["step"] + 1), _f(r["t"] / 60.0, decimal, 6),
            _f((r["t"] - step_start[r["step"]]) / 60.0, decimal, 6), _f(q, decimal), _f(0.0, decimal),
            _f(abs(r["i_ma"]) * 1e-3, decimal), _f(r["v"], decimal, 6), state[r["kind"]], "0",
            "02/01/2026 09:00:00",
        ]))
    path.write_text("\n".join(out) + "\n", encoding="cp1252", newline="\n")
    return path
