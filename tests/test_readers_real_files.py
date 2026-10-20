"""Real instrument exports, read end to end and checked against the files themselves (P08d).

Every file under data/public/<vendor>/ that this version can read is opened through
workflow.open_file without naming a profile, so recognition is part of what is tested. The
expected numbers are not taken from the reader: this module parses each file again with the csv
module or openpyxl, applies the unit and the sign it knows for that dialect, and computes the
quantities itself with numpy. The checks are physical where they can be:

- the arrays the reader produced equal the ones parsed here (units, decimal mark, sign);
- the elapsed test time equals the elapsed wall clock written in the same file, which fixes the
  time unit independently of any profile (Maccor's TestTime in minutes, Neware's Total Time
  across 24 h);
- over every step that carries current of one sign, the voltage rises exactly when the
  charge-positive current is positive (charging raises the terminal voltage);
- the instrument's own capacity counter agrees with the trapezoidal integral of the current;
- the engine's branch capacities equal the independent totals, and the number of charge and
  discharge branches equals the number of runs of one current sign found here;
- where the export carries the instrument's own cycle statistics (Landt's Cycle sheet), the
  engine's capacities agree with them within the file's printed resolution.

The sources and licences of the files are in each directory's ATTRIBUTION.md and manifest.json.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import numpy as np
import pytest

from electrochem import workflow
from electrochem.cycling import Settings
from electrochem.readers import BUILTIN, ImportFailure, Registry
from electrochem.readers.reader import read

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "data" / "public"


@pytest.fixture(autouse=True)
def _no_saved_profiles(tmp_path, monkeypatch):
    """Recognition must come from the built-in profiles, not from a profile saved on this machine."""
    monkeypatch.setenv("ELECTROCHEM_PROFILES", str(tmp_path / "profiles"))


# ------------------------------------------------------------ independent parsing
_CLOCK = re.compile(r"^\s*(?:(\d+)d\s*)?(\d+):(\d+):(\d+(?:\.\d*)?)\s*$")


def clock_s(tok: str) -> float:
    m = _CLOCK.match(tok)
    assert m, f"not a clock duration: {tok!r}"
    d, h, mi, s = m.groups()
    return int(d or 0) * 86400 + int(h) * 3600 + int(mi) * 60 + float(s)


def num(tok: str, decimal: str = ".") -> float:
    tok = tok.strip()
    if decimal == ",":
        tok = tok.replace(",", ".")
    return float(tok) if tok else float("nan")


def text_table(path: Path, starts_with: str, delimiter: str, encoding: str = "cp1252"):
    """(header, rows) of the table whose header line begins with `starts_with`; blank lines dropped."""
    lines = path.read_bytes().decode(encoding).splitlines()
    k = next(j for j, ln in enumerate(lines) if ln.startswith(starts_with))
    grid = list(csv.reader(io.StringIO("\n".join(lines[k:])), delimiter=delimiter))
    header = [h.strip().lstrip("﻿") for h in grid[0]]
    rows = [r for r in grid[1:] if any(c.strip() for c in r)]
    return header, rows


def excel_seconds(v) -> float:
    """An Excel time cell as elapsed seconds, from the serial openpyxl decoded it from."""
    from openpyxl.utils.datetime import to_excel
    if isinstance(v, dt.timedelta):
        return v.total_seconds()
    return float(to_excel(v)) * 86400.0


def xlsx_table(path: Path, first_name: str):
    """(header, rows) of the worksheet whose first row starts with `first_name`, as Python values."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            head = next(it, None)
            if head and head[0] == first_name:
                header = [str(h) for h in head]
                return header, [list(r) for r in it if any(c is not None for c in r)]
    finally:
        wb.close()
    raise AssertionError(f"no sheet starting with {first_name!r}")


def wall(fmt: str):
    """Wall-clock text as seconds, read as written (no time zone is applied, so the machine running
    the tests cannot shift it)."""
    epoch = dt.datetime(1970, 1, 1)

    def f(tok):
        tok = tok if isinstance(tok, dt.datetime) else str(tok).strip()
        try:
            d = tok if isinstance(tok, dt.datetime) else dt.datetime.strptime(tok, fmt)
        except ValueError:
            return float("nan")
        return (d - epoch).total_seconds()
    return f


@dataclass
class Spec:
    """How this module reads one file on its own, and what it expects of it."""
    path: str
    profile: str
    load: Callable                     # path -> (header, rows)
    t: Callable                        # (header, rows) -> seconds
    i: Callable                        # -> charge-positive A
    e: Callable                        # -> V
    step: Callable | None              # -> step key per row (None: runs of one current sign)
    wall: Callable | None = None       # -> wall-clock seconds per row
    counter: Callable | None = None    # -> signed per-row increments (Ah) of the instrument's counter
    counter_tol: float = 1e-3          # counter against trapezoid, each direction
    uses_counter: bool = True          # does the profile hand the counter to the engine
    compare_counter: bool = True       # is the counter-against-integral comparison meaningful here
    wall_tol_s: float = 2.0
    extra: dict = field(default_factory=dict)
    votes: int = 0                     # voting steps expected in the voltage-trend test (docs/P08d, S n/m)
    consistent: tuple | None = None    # counter behaviours the file leaves open (docs/P08d, C)


def col(h, rows, name, f=float):
    k = h.index(name)
    return np.array([f(r[k]) if k < len(r) else np.nan for r in rows], dtype=float)


def text_col(h, rows, name):
    k = h.index(name)
    return [str(r[k]).strip() if k < len(r) and r[k] is not None else "" for r in rows]


def by_state(h, rows, cur, state, sign):
    """Unsigned current made charge-positive from a state column."""
    s = np.array([sign.get(v, np.nan) for v in text_col(h, rows, state)])
    return np.abs(cur) * s


def _restarting_increments(q, group):
    """Per-row increase of a counter that restarts at each change of `group`: on the first row of a
    group the value itself (what accrued since the restart), elsewhere the difference. A fall
    inside a group is not charge, so it counts as zero."""
    inc = np.diff(q, prepend=q[0])
    starts = np.flatnonzero(np.diff(group, prepend=group[0]) != 0)
    inc[starts] = q[starts]
    inc[0] = 0.0
    return np.where(np.isfinite(inc) & (inc > 0), inc, 0.0)


def throughput_increments(q, group, i):
    """Signed per-row increments (Ah, charge positive) of one unsigned counter: the direction of
    each increment is the sign of the current on its row (or the row before, where that is 0)."""
    s = np.sign(np.nan_to_num(i))
    s = np.where(s == 0, np.concatenate(([0.0], s[:-1])), s)
    return _restarting_increments(q, group) * s


def two_counter_increments(qc, qd, group):
    """Signed per-row increments (Ah) from a charge counter and a discharge counter. The direction
    comes from which counter moved, not from the current, so the sign is checked as well."""
    return _restarting_increments(qc, group) - _restarting_increments(qd, group)


# ------------------------------------------------------------------ the files
ARBIN_RES = Spec(
    "arbin/arbin_example.csv", "arbin-res-table-csv",
    load=lambda p: text_table(p, "Data_Point", ",", "utf-8"),
    t=lambda h, r: col(h, r, "Test_Time"), i=lambda h, r: col(h, r, "Current"),
    e=lambda h, r: col(h, r, "Voltage"), step=lambda h, r: col(h, r, "Step_Index"),
    wall=lambda h, r: col(h, r, "DateTime"),
    counter=lambda h, r: two_counter_increments(col(h, r, "Charge_Capacity"), col(h, r, "Discharge_Capacity"),
                                                col(h, r, "Cycle_Index")),
    counter_tol=5e-3, wall_tol_s=2.0)

ARBIN_MITS8 = Spec(
    "arbin/sample_data_arbin.csv", "arbin-mits8-csv",
    load=lambda p: text_table(p, "﻿Data Point", ",", "utf-8"),
    t=lambda h, r: col(h, r, "Test Time (s)"), i=lambda h, r: col(h, r, "Current (A)"),
    e=lambda h, r: col(h, r, "Voltage (V)"), step=lambda h, r: col(h, r, "Step Index"),
    wall=lambda h, r: col(h, r, "Date Time", wall("%m/%d/%Y %H:%M:%S.%f")),
    counter=lambda h, r: two_counter_increments(col(h, r, "Charge Capacity (Ah)"),
                                                col(h, r, "Discharge Capacity (Ah)"), col(h, r, "Cycle Index")),
    counter_tol=5e-3, wall_tol_s=0.01, extra={"no_vote": "two charging rows only"})

ARBIN_MITS = Spec(
    "arbin/arbin.csv", "arbin-csv",
    load=lambda p: text_table(p, "Data_Point", ",", "utf-8"),
    t=lambda h, r: col(h, r, "Test_Time(s)"), i=lambda h, r: col(h, r, "Current(A)"),
    e=lambda h, r: col(h, r, "Voltage(V)"), step=lambda h, r: col(h, r, "Step_Index"),
    counter=lambda h, r: two_counter_increments(col(h, r, "Charge_Capacity(Ah)"),
                                                col(h, r, "Discharge_Capacity(Ah)"), col(h, r, "Cycle_Index")),
    compare_counter=False, extra={"no_vote": "one rest row repeated"})

ARBIN_XLSX = Spec(
    "arbin/bs542_004_gr_li_50ua_50mv_1v_191020_Channel_11.xlsx", "arbin-xlsx",
    load=lambda p: xlsx_table(p, "Date_Time"),
    t=lambda h, r: col(h, r, "Test_Time(s)"), i=lambda h, r: col(h, r, "Current(A)"),
    e=lambda h, r: col(h, r, "Voltage(V)"), step=lambda h, r: col(h, r, "Step_Index"),
    wall=lambda h, r: col(h, r, "Date_Time", wall("")),
    counter=lambda h, r: two_counter_increments(col(h, r, "Charge_Capacity(Ah)"),
                                                col(h, r, "Discharge_Capacity(Ah)"), col(h, r, "Cycle_Index")),
    counter_tol=1e-3, wall_tol_s=2.0)

MACCOR_STATE = {"C": 1.0, "D": -1.0, "R": 0.0, "O": 0.0}


def _maccor_tri(path, wall_fmt="%m/%d/%Y %H:%M:%S"):
    return Spec(
        path, "maccor-export-sec",
        load=lambda p: text_table(p, "Rec#", "\t"),
        t=lambda h, r: col(h, r, "Test (Sec)"), i=lambda h, r: col(h, r, "Amps"),
        e=lambda h, r: col(h, r, "Volts"), step=lambda h, r: col(h, r, "Step"),
        wall=lambda h, r: col(h, r, "DPt Time", wall(wall_fmt)),
        counter=lambda h, r: throughput_increments(col(h, r, "Amp-hr"), col(h, r, "Step"), col(h, r, "Amps")),
        counter_tol=1e-4)


MACCOR_TRI = _maccor_tri("maccor/xTESLADIAG_000038con.078")
MACCOR_TRI_MIDTEST = _maccor_tri("maccor/PredictionDiagnostics_000151_test.052")

MACCOR_UBHAM = Spec(
    "maccor/maccor_001_first_cycle.txt", "maccor-export-clock",
    load=lambda p: text_table(p, "Rec#", "\t"),
    t=lambda h, r: col(h, r, "TestTime", clock_s),
    i=lambda h, r: by_state(h, r, col(h, r, "Amps"), "State", MACCOR_STATE),
    e=lambda h, r: col(h, r, "Volts"), step=lambda h, r: col(h, r, "Step"),
    wall=lambda h, r: col(h, r, "DPt Time", wall("%m/%d/%Y %H:%M:%S")),
    counter=lambda h, r: throughput_increments(
        col(h, r, "Amp-hr"), col(h, r, "Step"), by_state(h, r, col(h, r, "Amps"), "State", MACCOR_STATE)),
    counter_tol=1e-4)

MACCOR_KIT = Spec(
    "maccor/maccor_002.txt", "maccor-txt",
    load=lambda p: text_table(p, "Rec\tCycle P", "\t"),
    t=lambda h, r: col(h, r, "TestTime") * 60.0,             # minutes: see the wall-clock check
    i=lambda h, r: by_state(h, r, col(h, r, "Current [A]"), "Md", MACCOR_STATE),
    e=lambda h, r: col(h, r, "Voltage [V]"), step=lambda h, r: col(h, r, "Step"),
    wall=lambda h, r: col(h, r, "DPT Time", wall("%m/%d/%Y %I:%M:%S %p")),
    counter=lambda h, r: throughput_increments(
        col(h, r, "Cap. [Ah]"), col(h, r, "Step"), by_state(h, r, col(h, r, "Current [A]"), "Md", MACCOR_STATE)),
    counter_tol=1e-4)

NEWARE_CSV = Spec(
    "neware/neware_uio.csv", "neware-csv",
    load=lambda p: text_table(p, "DataPoint", ",", "utf-8"),
    t=lambda h, r: col(h, r, "Cumulative Time", clock_s), i=lambda h, r: col(h, r, "Current(A)"),
    e=lambda h, r: col(h, r, "Voltage(V)"), step=lambda h, r: col(h, r, "Step Index"),
    wall=lambda h, r: col(h, r, "Date", wall("%Y-%m-%d %H:%M:%S")),
    counter=lambda h, r: two_counter_increments(col(h, r, "Chg. Cap.(Ah)"), col(h, r, "DChg. Cap.(Ah)"),
                                                col(h, r, "Step Index")),
    counter_tol=1e-3)


def _neware_xlsx(path):
    return Spec(
        path, "neware-xlsx",
        load=lambda p: xlsx_table(p, "Date"),
        t=lambda h, r: col(h, r, "Total Time", excel_seconds),
        i=lambda h, r: col(h, r, "Current(mA)") * 1e-3,
        e=lambda h, r: col(h, r, "Voltage(V)"), step=lambda h, r: col(h, r, "Step Index"),
        wall=lambda h, r: col(h, r, "Date", wall("")),
        counter=lambda h, r: two_counter_increments(col(h, r, "Chg. Cap.(mAh)") * 1e-3,
                                                    col(h, r, "DChg. Cap.(mAh)") * 1e-3, col(h, r, "Step Index")),
        counter_tol=1e-4, wall_tol_s=1.0)


NEWARE_XLSX_START = _neware_xlsx("neware/sample_data_neware_rows_2_25001.xlsx")
NEWARE_XLSX_24H = _neware_xlsx("neware/sample_data_neware_rows_80001_90000.xlsx")

BASYTEC = Spec(
    "basytec/sample_data_basytec.txt", "basytec-txt",
    load=lambda p: text_table(p, "~Time[s]", "\t", "utf-8"),
    t=lambda h, r: col(h, r, "~Time[s]"), i=lambda h, r: col(h, r, "I[A]"),
    e=lambda h, r: col(h, r, "U[V]"), step=lambda h, r: col(h, r, "Line"),
    counter=lambda h, r: throughput_increments(col(h, r, "Ah-Step"), col(h, r, "Line"), col(h, r, "I[A]")),
    counter_tol=1e-3, uses_counter=False)

NOVONIX = Spec(
    "novonix/sample_data_novonix.csv", "novonix-csv",
    load=lambda p: text_table(p, "Date and Time,", ",", "utf-8"),
    t=lambda h, r: col(h, r, "Run Time (h)") * 3600.0, i=lambda h, r: col(h, r, "Current (A)"),
    e=lambda h, r: col(h, r, "Potential (V)"), step=lambda h, r: col(h, r, "Step Number"),
    wall=lambda h, r: col(h, r, "Date and Time", wall("%Y-%m-%d %H:%M:%S")),
    counter=lambda h, r: throughput_increments(col(h, r, "Capacity (Ah)"), col(h, r, "Step Number"),
                                               col(h, r, "Current (A)")),
    counter_tol=1e-3, uses_counter=False)

PEC = Spec(
    "pec/pec_first_cycle.csv", "pec-csv",
    load=lambda p: text_table(p, "Test,Cell", ","),
    t=lambda h, r: col(h, r, "Total Time (Seconds)"), i=lambda h, r: col(h, r, "Current (mA)") * 1e-3,
    e=lambda h, r: col(h, r, "Voltage (mV)") * 1e-3, step=lambda h, r: col(h, r, "Step"),
    wall=lambda h, r: col(h, r, "Real Time", wall("%m/%d/%Y %H:%M:%S")),
    counter=lambda h, r: two_counter_increments(col(h, r, "Charge Capacity (mAh)") * 1e-3,
                                                col(h, r, "Discharge Capacity (mAh)") * 1e-3, col(h, r, "Cycle")),
    counter_tol=1e-4)

PEC_IDLE = Spec(
    "pec/Test25205.csv", "pec-csv-clock",
    load=lambda p: text_table(p, "Test,Cell", ","),
    t=lambda h, r: col(h, r, "Total Time (Hours in hh:mm:ss.xxx)", clock_s),
    i=lambda h, r: col(h, r, "Current (A)"), e=lambda h, r: col(h, r, "Voltage (V)"),
    step=lambda h, r: col(h, r, "Step"),
    wall=lambda h, r: col(h, r, "Real Time", wall("%m/%d/%Y %H:%M:%S")),
    counter=lambda h, r: two_counter_increments(col(h, r, "Charge Capacity (mAh)") * 1e-3,
                                                col(h, r, "Discharge Capacity (mAh)") * 1e-3, col(h, r, "Cycle")),
    compare_counter=False, wall_tol_s=1.5, extra={"no_vote": "open circuit"})

LANDT_STATE = {"C_CC": 1.0, "D_CC": -1.0, "R": 0.0}
LANDT = Spec(
    "landt/NJK_CC_156_C30_C_x30.xlsx", "landt-xlsx",
    load=lambda p: xlsx_table(p, "Index"),
    t=lambda h, r: col(h, r, "TestTime/h") * 3600.0, i=lambda h, r: col(h, r, "Current/mA") * 1e-3,
    e=lambda h, r: col(h, r, "Voltage/V"), step=None, counter=None, uses_counter=False)

_EXPECT = {  # spec: (voting steps, counter behaviours left open by the file)
    "ARBIN_RES": (6, ("cycle_cumulative",)), "ARBIN_MITS8": (0, ("cycle_cumulative", "step_cumulative", "test_cumulative")), "ARBIN_MITS": (0, ("cycle_cumulative", "step_cumulative", "test_cumulative")),
    "ARBIN_XLSX": (3, ("cycle_cumulative", "test_cumulative")), "MACCOR_TRI": (16, ("step_cumulative",)),
    "MACCOR_TRI_MIDTEST": (1, ("cycle_cumulative", "step_cumulative", "test_cumulative")), "MACCOR_UBHAM": (3, ("step_cumulative",)),
    "MACCOR_KIT": (4, ("step_cumulative",)), "NEWARE_CSV": (16, ("step_cumulative",)),
    "NEWARE_XLSX_START": (1, ("step_cumulative",)), "NEWARE_XLSX_24H": (1, ("cycle_cumulative", "step_cumulative", "test_cumulative")),
    "BASYTEC": (1, None), "NOVONIX": (1, None), "PEC": (3, ("cycle_cumulative",)), "PEC_IDLE": (0, ("cycle_cumulative", "step_cumulative", "test_cumulative")),
    "LANDT": (60, None),
}
for _name, (_votes, _open) in _EXPECT.items():
    globals()[_name].votes, globals()[_name].consistent = _votes, _open

ANALYSED = [ARBIN_RES, ARBIN_MITS8, ARBIN_MITS, ARBIN_XLSX, MACCOR_TRI, MACCOR_TRI_MIDTEST, MACCOR_UBHAM,
            MACCOR_KIT, NEWARE_CSV, NEWARE_XLSX_START, NEWARE_XLSX_24H, BASYTEC, NOVONIX, PEC, PEC_IDLE, LANDT]


# ------------------------------------------------------------- shared checks
def steps_of(spec, h, rows, i):
    if spec.step is not None:
        return spec.step(h, rows)
    sgn = np.sign(np.nan_to_num(i))
    return np.cumsum(np.concatenate(([0], sgn[1:] != sgn[:-1])))


def runs(key):
    starts = np.flatnonzero(np.concatenate(([True], key[1:] != key[:-1])))
    return list(zip(starts, np.concatenate((starts[1:], [key.size]))))


def trapezoid_edges(t, i):
    """Signed charge (Ah) on each edge k-1 -> k by the trapezoidal rule; 0 where undefined."""
    h = np.diff(t)
    ok = np.isfinite(h) & (h > 0) & np.isfinite(i[1:]) & np.isfinite(i[:-1])
    return np.concatenate(([0.0], np.where(ok, h * (i[1:] + i[:-1]) / 2.0 / 3600.0, 0.0)))


def by_direction(q, mask):
    q = np.where(mask, q, 0.0)
    return float(q[q > 0].sum()), float(-q[q < 0].sum())


def within(key):
    """Edges k-1 -> k whose two rows share `key` (a step, or a branch)."""
    return np.concatenate(([False], key[1:] == key[:-1]))


def branch_key(i):
    """A label per row for its run of one current sign; rows without current get a label of their own."""
    s = np.sign(np.nan_to_num(i))
    run = np.cumsum(np.concatenate(([0], s[1:] != s[:-1])))
    return np.where(s == 0, -1 - np.arange(s.size), run)


def sign_runs(i):
    """Maximal runs of rows carrying current of one sign: what a branch is."""
    s = np.sign(np.nan_to_num(i))
    out = []
    k = 0
    while k < s.size:
        if s[k] == 0:
            k += 1
            continue
        j = k
        while j + 1 < s.size and s[j + 1] == s[k]:
            j += 1
        out.append(int(s[k]))
        k = j + 1
    return out


@pytest.fixture(scope="module")
def opened():
    """Each file opened once, by recognition, and analysed with the settings of the workflow tests."""
    import os
    import tempfile
    cache = {}
    old = os.environ.get("ELECTROCHEM_PROFILES")
    os.environ["ELECTROCHEM_PROFILES"] = tempfile.mkdtemp()
    try:
        for spec in ANALYSED:
            p = PUBLIC / spec.path
            L = workflow.open_file(p)
            run = workflow.analyse_file(L, Settings(charge="ox", terminal_voltage=True))
            h, rows = spec.load(p)
            cache[spec.path] = (L, run, h, rows)
    finally:
        if old is None:
            os.environ.pop("ELECTROCHEM_PROFILES", None)
        else:
            os.environ["ELECTROCHEM_PROFILES"] = old
    return cache


ids = [s.path for s in ANALYSED]


@pytest.mark.parametrize("spec", ANALYSED, ids=ids)
def test_the_file_is_recognised_by_its_columns_alone(spec):
    """No profile is named: the registry finds exactly one built-in profile for the file."""
    p = PUBLIC / spec.path
    from electrochem.readers.propose import detect_layout
    from electrochem.readers.registry import claims, sniff_text
    layout, headers, _ = detect_layout(p)
    text = sniff_text(p, headers)
    claimed = [q.name for q in BUILTIN.values() if claims(q, p, layout.kind, headers, text)]
    assert claimed == [spec.profile]


@pytest.mark.parametrize("spec", ANALYSED, ids=ids)
def test_the_arrays_equal_the_file_parsed_independently(opened, spec):
    """Units, decimal mark, clock formats and sign, checked against a parse made here."""
    L, run, h, rows = opened[spec.path]
    assert L.proposal is None and L.meta["profile"] == spec.profile
    tab = L.tab
    assert tab.n == len(rows)
    t, i, e = spec.t(h, rows), spec.i(h, rows), spec.e(h, rows)
    assert np.allclose(tab.t, t, rtol=0, atol=1e-6, equal_nan=True)
    assert np.allclose(tab.i, i, rtol=1e-12, atol=1e-15, equal_nan=True)
    assert np.allclose(tab.e, e, rtol=1e-12, atol=1e-12, equal_nan=True)
    assert np.all(np.diff(t[np.isfinite(t)]) >= 0), "test time must never run backwards"


@pytest.mark.parametrize("spec", [s for s in ANALYSED if s.wall], ids=[s.path for s in ANALYSED if s.wall])
def test_the_time_unit_is_confirmed_by_the_wall_clock_in_the_same_file(opened, spec):
    """Elapsed test time from the reader equals elapsed wall-clock time written in the file.

    This is what shows that Maccor's TestTime in the 12-line-header export is in minutes, that
    Neware's Total Time is continuous across the 24 h at which Excel stops writing it as a time of
    day, and that the clock formats are read as hours, minutes and seconds. The tolerance, 2 s plus
    5e-4 of the span, admits the drift seen between an instrument's timer and its PC's clock
    (2e-4 in the Neware and TRI Maccor files) and rejects any wrong unit by a wide margin. The wall
    clock is read as written, with no time zone, so the machine running the tests cannot move it."""
    L, run, h, rows = opened[spec.path]
    w = spec.wall(h, rows)
    ok = np.isfinite(w) & np.isfinite(L.tab.t)
    assert ok.sum() > 1
    t, w = L.tab.t[ok], w[ok]
    elapsed_test = t[-1] - t[0]
    elapsed_wall = w[-1] - w[0]
    assert elapsed_test > 0
    assert abs(elapsed_test - elapsed_wall) <= spec.wall_tol_s + 5e-4 * elapsed_wall


@pytest.mark.parametrize("spec", ANALYSED, ids=ids)
def test_charging_raises_the_voltage_in_every_step(opened, spec):
    """The sign convention, measured: over each step of five or more rows that carries current of
    one sign, the voltage ends higher than it began exactly when the charge-positive current is
    positive. Independent of the reader's own check, which it also has to agree with."""
    L, run, h, rows = opened[spec.path]
    i, e = spec.i(h, rows), spec.e(h, rows)
    step = steps_of(spec, h, rows, i)
    votes = []
    for a, b in runs(step):
        ii, ee = i[a:b], e[a:b]
        flowing = np.isfinite(ii) & np.isfinite(ee) & (ii != 0)        # a step's zero-current rows abstain
        ii, ee = ii[flowing], ee[flowing]
        if ii.size < 5 or not (np.all(ii > 0) or np.all(ii < 0)) or abs(ee[-1] - ee[0]) < 1e-3:
            continue
        votes.append(np.sign(ee[-1] - ee[0]) == np.sign(ii[0]))
    assert all(votes)
    assert len(votes) == spec.votes
    assert votes or spec.extra.get("no_vote"), "a file with current but no voting step tests nothing here"
    rep = L.meta["import_report"]["sign_vs_voltage"]
    assert rep["disagree"] == 0
    if spec.step is not None:
        assert rep["agree"] == len(votes)


@pytest.mark.parametrize("spec", ANALYSED, ids=ids)
def test_the_counter_behaviour_is_measured_on_the_file(opened, spec):
    """Which reset behaviours (test-, cycle-, step-cumulative) the file leaves open, as the reader
    measured them: one behaviour means the declaration is verified on this file, several mean the
    file cannot decide (one cycle, one step), none bound means the profile integrates."""
    L, run, h, rows = opened[spec.path]
    rep = L.meta["import_report"]
    got = rep["capacity_semantics_consistent"]
    assert (tuple(got) if got is not None else None) == spec.consistent
    assert rep["capacity_semantics_verified"] == (spec.consistent is not None and len(spec.consistent) == 1)


COUNTED = [s for s in ANALYSED if s.counter and s.compare_counter]


@pytest.mark.parametrize("spec", COUNTED, ids=[s.path for s in COUNTED])
def test_the_instrument_counter_agrees_with_the_integral_of_the_current(opened, spec):
    """Two independent statements of the same charge, both computed here from the raw file: the
    instrument's counter and the trapezoidal integral of the sampled current, over the edges that
    lie inside one programmed step (at a step boundary the instrument counts charge it did not
    log a current for). Charge and discharge are compared separately."""
    L, run, h, rows = opened[spec.path]
    t, i = spec.t(h, rows), spec.i(h, rows)
    inside = within(steps_of(spec, h, rows, i))
    qc, qd = by_direction(spec.counter(h, rows), inside)
    tc, td = by_direction(trapezoid_edges(t, i), inside)
    total = qc + qd
    assert total > 0
    for counter, integral in ((qc, tc), (qd, td)):
        assert abs(counter - integral) <= spec.counter_tol * max(counter, integral, 0.1 * total)


@pytest.mark.parametrize("spec", ANALYSED, ids=ids)
def test_the_engine_capacities_equal_the_independent_totals(opened, spec):
    """End to end: the engine's branch capacities, summed by direction, equal what this module
    computes from the file over the edges inside each run of one current sign, from the counter
    where the profile hands it over and from the trapezoid where it does not; and there is one
    branch for each such run."""
    L, run, h, rows = opened[spec.path]
    t, i = spec.t(h, rows), spec.i(h, rows)
    br = run.an.branches
    got = {d: sum(b.charge_mAh for b in br if b.direction == d) / 1e3 for d in ("charge", "discharge")}
    # the engine does not integrate an edge of zero duration and drops the counter's increment on
    # it too: Neware's CSV clock has 1 s resolution, so sub-second records share a timestamp
    inside = within(branch_key(i)) & (np.diff(t, prepend=t[0]) > 0)
    if spec.uses_counter and spec.counter:
        assert L.meta["import_report"]["instrument_increment_available"]
        want_c, want_d = by_direction(spec.counter(h, rows), inside)
        assert {b.charge_method for b in br} <= {"instrument dq increments"}
    else:
        assert not L.meta["import_report"]["instrument_increment_available"]
        want_c, want_d = by_direction(trapezoid_edges(t, i), inside)
        assert {b.charge_method for b in br} <= {"trapezoid of sampled current"}
    scale = max(want_c + want_d, 1e-30)
    assert got["charge"] == pytest.approx(want_c, rel=1e-6, abs=1e-9 * scale)
    assert got["discharge"] == pytest.approx(want_d, rel=1e-6, abs=1e-9 * scale)
    directions = [1 if b.direction == "charge" else -1 for b in br]
    assert directions == sign_runs(i)


# ------------------------------------------------------- file-specific checks
def test_landt_capacities_agree_with_the_instruments_own_cycle_sheet(opened):
    """The Landt export carries the instrument's per-cycle charge and discharge capacities (Cycle
    sheet, printed to 0.0001 mAh). This cell starts with a discharge, so the instrument's cycle k
    is the k-th discharge followed by the k-th charge. The record sheet's time is printed to
    0.01 h, which with the printed capacity bounds the agreement to about 1.5 %."""
    L, run, h, rows = opened[LANDT.path]
    from openpyxl import load_workbook
    wb = load_workbook(PUBLIC / LANDT.path, read_only=True, data_only=True)
    sheet = [list(r) for r in wb["Cycle"].iter_rows(values_only=True)]
    wb.close()
    head, body = sheet[0], [r for r in sheet[1:] if r[0] and r[0] > 0]
    vend_c = [r[head.index("Charge-Cap/mAh")] for r in body]
    vend_d = [r[head.index("Discharge-Cap/mAh")] for r in body]
    ours_c = [b.charge_mAh for b in run.an.branches if b.direction == "charge"]
    ours_d = [b.charge_mAh for b in run.an.branches if b.direction == "discharge"]
    assert len(vend_c) == len(vend_d) == 30 and len(ours_c) == len(ours_d) == 30
    for k, (a, b) in enumerate(zip(ours_c, vend_c)):
        assert a == pytest.approx(b, rel=0.025), f"charge, cycle {k + 1}"
    for k, (a, b) in enumerate(zip(ours_d, vend_d)):
        assert a == pytest.approx(b, rel=0.025), f"discharge, cycle {k + 1}"


def test_neware_total_time_is_continuous_where_excel_switches_to_a_date(opened):
    """In the 24 h window the raw cells change type (time of day, then a 1900 date) at the row
    where Total Time passes one day; the elapsed seconds must step by the sampling interval there,
    not by a day or by zero."""
    from openpyxl import load_workbook
    p = PUBLIC / NEWARE_XLSX_24H.path
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    kinds = [type(r[2]).__name__ for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    switch = kinds.index("datetime")
    assert set(kinds[:switch]) == {"time"} and set(kinds[switch:]) == {"datetime"}
    L, run, h, rows = opened[NEWARE_XLSX_24H.path]
    t = L.tab.t
    assert t[switch - 1] < 86400.0 <= t[switch]
    assert 0 < t[switch] - t[switch - 1] < 2.0
    assert np.max(np.diff(t)) < 2.0


def test_maccor_kit_time_is_minutes_and_the_counter_needs_them(opened):
    """With the time read as seconds, the integral of the current would be 60 times smaller than
    the instrument's counter: the unit is fixed by the counter as well as by the wall clock."""
    L, run, h, rows = opened[MACCOR_KIT.path]
    t_as_s = col(h, rows, "TestTime")
    i = MACCOR_KIT.i(h, rows)
    inside = within(col(h, rows, "Step"))
    tc_s, td_s = by_direction(trapezoid_edges(t_as_s, i), inside)
    qc, qd = by_direction(MACCOR_KIT.counter(h, rows), inside)
    assert qc / tc_s == pytest.approx(60.0, rel=1e-3) and qd / td_s == pytest.approx(60.0, rel=1e-3)


def test_a_file_that_starts_part_way_through_a_test_keeps_its_counters_honest(opened):
    """beep's PredictionDiagnostics file begins at record 249 791, in cycle 37, in the middle of a
    discharge step whose Amp-hr already reads 0.0192 Ah. Only what was discharged inside the file
    is counted, and the cycle labels carry the instrument's numbering (warned, not renumbered)."""
    L, run, h, rows = opened[MACCOR_TRI_MIDTEST.path]
    q = col(h, rows, "Amp-hr")
    assert q[0] == pytest.approx(0.0191579754)
    (b,) = run.an.branches
    assert b.direction == "discharge"
    assert b.charge_mAh / 1e3 == pytest.approx(q[-1] - q[0], rel=1e-9)
    assert any(f["code"] == "cycle_base_mismatch" for f in L.meta["import_findings"])


def test_the_arbin_res_table_counter_resets_per_cycle_and_is_measured_so(opened):
    L, run, h, rows = opened[ARBIN_RES.path]
    rep = L.meta["import_report"]
    assert rep["capacity_semantics_consistent"] == ["cycle_cumulative"] and rep["capacity_semantics_verified"]
    qc, cyc, stp = col(h, rows, "Charge_Capacity"), col(h, rows, "Cycle_Index"), col(h, rows, "Step_Index")
    k = np.flatnonzero(np.diff(cyc) != 0)[0] + 1                    # first cycle boundary
    assert qc[k - 1] > 1.0 and qc[k] < 1e-6                          # resets there
    within = np.flatnonzero((np.diff(stp) != 0) & (np.diff(cyc) == 0)) + 1
    assert np.any(qc[within] >= qc[within - 1] - 1e-12)              # carried across a step boundary


@pytest.mark.parametrize("path,profile", [("maccor/maccor_example.001", "maccor-export-min"),
                                          ("maccor/maccor_two.txt", "maccor-txt")])
def test_a_rest_only_file_is_recognised_but_its_sign_cannot_be_established(path, profile, tmp_path):
    """Both files hold nothing but rest: the profile is found, and the import stops because no row
    shows which sign is charge. A person's answer (the current is unsigned) then reads it, and the
    time unit still matches the wall clock written beside it."""
    p = PUBLIC / path
    assert Registry(tmp_path / "none").identify(p, list(BUILTIN.values())).name == profile
    with pytest.raises(ImportFailure, match="ambiguous_sign"):
        workflow.open_file(p)
    prof = BUILTIN[profile]
    answered = prof.model_copy(update={"conventions": prof.conventions.model_copy(update={"current_sign": "unsigned"})})
    raw = read(p, answered)
    if profile == "maccor-txt":
        h, rows = text_table(p, "Rec\tCycle P", "\t")
        assert raw.report["decimal"] == ","
        w = col(h, rows, "DPT Time", wall("%m/%d/%Y %H:%M:%S"))
        tm = col(h, rows, "TestTime", lambda s: num(s, ","))
        assert np.allclose(raw.t, tm * 60.0, equal_nan=True)
        # the last line was cut off mid-number when the export was taken: kept, without a step type
        assert any(f["code"] == "step_type_missing" for f in raw.findings)
    else:
        h, rows = text_table(p, "Rec#", "\t")
        w = col(h, rows, "DPt Time", wall("%m/%d/%Y %H:%M:%S"))
    ok = np.isfinite(w) & np.isfinite(raw.t)
    a, b = np.flatnonzero(ok)[[0, -1]]
    assert abs((raw.t[b] - raw.t[a]) - (w[b] - w[a])) <= 1.0


def test_an_open_circuit_file_is_read_and_its_counter_is_not_compared(opened):
    """PEC's 2026 export of an idle cell: the counter creeps by 0.02 mAh of leakage that the logged
    current (0 or a few uA) cannot show, so the comparison is not made rather than failed."""
    L, run, h, rows = opened[PEC_IDLE.path]
    codes = {f["code"] for f in L.meta["import_findings"]}
    assert "counter_not_compared" in codes and "counter_disagrees_with_integral" not in codes
    assert not [b for b in run.an.branches if b.duration_s > 60]


def test_the_header_only_arbin_csv_is_read_but_holds_no_branch(opened):
    """galv-harvester's arbin.csv repeats one rest row 21 times: the MITS Pro column names are
    real, the data say nothing about conventions."""
    L, run, h, rows = opened[ARBIN_MITS.path]
    assert len(rows) == 21 and len({tuple(r) for r in rows}) == 1
    assert run.an.branches == [] or all(b.charge_mAh == 0 for b in run.an.branches)


@pytest.mark.parametrize("path", ["arbin/arbin_example.res", "neware/test.nda", "neware/test.ndax"])
def test_a_binary_instrument_file_is_refused_with_a_reason(path):
    """Shipped as examples of what this version does not read (docs/P08d, binary formats)."""
    p = PUBLIC / path
    head = p.read_bytes()[:4096]
    assert b"\x00" in head or head[:2] == b"PK"                       # binary, or a zip container (.ndax)
    with pytest.raises(ImportFailure, match="binary_export"):
        workflow.open_file(p)


def test_every_shipped_file_is_listed_with_its_source_and_licence():
    """No file under the vendor directories without a manifest entry, an ATTRIBUTION.md and the
    licence text it names; and the shipped bytes are the ones the manifest records."""
    import hashlib
    import json
    vendors = ["arbin", "maccor", "neware", "basytec", "novonix", "pec", "landt"]
    for v in vendors:
        d = PUBLIC / v
        entries = json.loads((d / "manifest.json").read_text())
        names = {e["file"] for e in entries}
        data = {f.name for f in d.iterdir()
                if f.is_file() and f.name not in ("manifest.json", "ATTRIBUTION.md") and not f.name.startswith("LICENSE")}
        assert data == names, v
        assert (d / "ATTRIBUTION.md").is_file()
        for e in entries:
            assert hashlib.sha256((d / e["file"]).read_bytes()).hexdigest() == e["sha256"]
            assert (d / e["license_file"]).is_file()
            assert e["license"] in ("MIT", "BSD-2-Clause", "BSD-3-Clause", "Apache-2.0")
            assert e["file"] in (d / "ATTRIBUTION.md").read_text(encoding="utf-8")


def test_maccor_signed_export_profile_on_a_synthetic_file_in_its_layout(tmp_path):
    """maccor-export-signed was measured on battetl's Maccor exports, which cannot be shipped
    (BattSoft Community License). This is NOT a real file: it copies their layout exactly (two
    preamble lines with Maccor's doubled apostrophe, a header ending in a tab, signed Current(A),
    one Capacity(Ah) counter restarting at every step, no State column) around a known experiment,
    so that the profile is at least exercised here. docs/P08d lists it as verified on inspected
    files only."""
    from tests import reader_fixtures as fx
    head = ["Cyc#", "Step", "TestTime(s)", "StepTime(s)", "Capacity(Ah)", "Watt-hr", "Current(A)",
            "Voltage(V)", "ES", "DPt Time", "ACR", "DCIR", "Temp 1"]
    out = ["Today''s Date\t02/27/2023 11:18:36 AM", "Date of Test:\t03/26/2020 11:00:14 PM", "\t".join(head) + "\t"]
    start = {}
    t0 = dt.datetime(2020, 3, 26, 23, 0, 14)
    for r in fx.ground_truth():
        start.setdefault(r["step"], r["t"])
        stamp = (t0 + dt.timedelta(seconds=r["t"])).strftime("%m/%d/%Y %I:%M:%S %p")
        out.append("\t".join([str(r["cycle"] - 1), str(r["step"] + 1), f"{r['t']:.2f}", f"{r['t'] - start[r['step']]:.2f}",
                              f"{r['sqc_ah'] + r['sqd_ah']:.9f}", "0.00", f"{r['i_ma'] * 1e-3:.9f}", f"{r['v']:.5f}",
                              "0", stamp, "0.000", "0.000", "25.00"]))
    p = tmp_path / "BG_layout - 079.txt"
    p.write_text("\n".join(out) + "\n", encoding="cp1252")
    L = workflow.open_file(p)
    assert L.meta["profile"] == "maccor-export-signed"
    run = workflow.analyse_file(L, Settings(charge="ox", terminal_voltage=True))
    q = {b.direction: b.charge_mAh for b in run.an.branches}
    assert q["charge"] == pytest.approx(1.0, rel=1e-9) and q["discharge"] == pytest.approx(0.5, rel=1e-9)
    assert L.meta["import_report"]["capacity_semantics_verified"]
